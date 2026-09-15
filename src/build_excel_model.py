"""
Construit excel/modele_arbitrage.xlsx : le classeur de decision.

Ce script est le "code source" du classeur - le fichier .xlsx n'est pas
cree/modifie a la main, il est genere par ce script, versionne comme
n'importe quel autre artefact reproductible. Pour le regenerer :

    python -m src.build_excel_model

Trois feuilles :
- "Hypotheses" : les valeurs d'entree, separees en deux blocs (voir plus
  bas pourquoi).
- "Netback" : la meme formule que src/netback.py, ecrite en formules
  Excel, colonne par destination.
- "Sensibilite" : la table a deux variables (taux de fret x ecart
  JKM-TTF) qui montre la frontiere de bascule Europe/Asie.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import config

CHEMIN_SORTIE = Path(__file__).resolve().parent.parent / "excel" / "modele_arbitrage.xlsx"

TITRE_FONT = Font(bold=True, size=13)
ENTETE_FONT = Font(bold=True)
ENTETE_FILL = PatternFill("solid", fgColor="1F4E78")
ENTETE_FONT_BLANC = Font(bold=True, color="FFFFFF")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")  # jaune pale : cellule modifiable


def _entete(ws, cellule, texte):
    c = ws[cellule]
    c.value = texte
    c.font = ENTETE_FONT_BLANC
    c.fill = ENTETE_FILL
    c.alignment = Alignment(horizontal="center")


def _construire_feuille_hypotheses(wb):
    ws = wb.active
    ws.title = "Hypotheses"
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 16

    ws["A1"] = "Hypotheses structurelles (source : src/config.py)"
    ws["A1"].font = TITRE_FONT

    # Ces valeurs viennent de config.py : stables sur plusieurs mois,
    # on ne s'attend pas a devoir les changer souvent (voir README,
    # "Fraicheur des hypotheses").
    structurelles = [
        ("Taille de cargaison (TBtu)", config.CARGO_SIZE_TBTU),
        ("Boil-off (%/jour)", config.BOIL_OFF_RATE_PER_DAY),
        ("Multiplicateur feedgas (x Henry Hub)", config.FEEDGAS_HH_MULTIPLIER),
        ("Tolling fee liquefaction ($/MMBtu)", config.LIQUEFACTION_TOLLING_FEE_USD_PER_MMBTU),
        ("Jours de mer - Europe", config.ROUTE_DAYS["europe"]),
        ("Jours de mer - Asie (Panama)", config.ROUTE_DAYS["asia_panama"]),
        ("Jours de mer - Asie (Cap)", config.ROUTE_DAYS["asia_cape"]),
        ("Peage du canal de Panama ($)", config.PANAMA_TOLL_USD),
        ("Conversion MWh -> MMBtu", config.MWH_TO_MMBTU),
        ("Frais de regaz - Europe ($/MMBtu)", 0.40),
        ("Frais de regaz - Asie ($/MMBtu)", 0.50),
    ]
    ligne = 3
    for libelle, valeur in structurelles:
        ws.cell(row=ligne, column=1, value=libelle)
        ws.cell(row=ligne, column=2, value=valeur)
        ligne += 1
    # lignes 3 a 13

    ligne_marche_titre = ligne + 1  # 15
    ws.cell(row=ligne_marche_titre, column=1,
            value="Parametres de marche (modifiables - a mettre a jour depuis l'actu)")
    ws.cell(row=ligne_marche_titre, column=1).font = TITRE_FONT

    ligne = ligne_marche_titre + 2  # 17
    marche = [
        ("TTF (EUR/MWh)", 32.00),
        ("Taux de change EUR/USD", 1.08),
        ("Henry Hub ($/MMBtu)", 3.20),
        ("JKM ($/MMBtu) - hypothese fragile, pas de source gratuite", 11.50),
        ("Taux d'affretement ($/jour) - hypothese fragile", 65_000),
    ]
    for libelle, valeur in marche:
        ws.cell(row=ligne, column=1, value=libelle)
        cell_valeur = ws.cell(row=ligne, column=2, value=valeur)
        cell_valeur.fill = INPUT_FILL
        ligne += 1
    # lignes 17 a 21

    ligne_route = ligne  # 22
    ws.cell(row=ligne_route, column=1, value="Route Asie retenue")
    cell_route = ws.cell(row=ligne_route, column=2, value="Panama")
    cell_route.fill = INPUT_FILL
    dv = DataValidation(type="list", formula1='"Panama,Cap"', allow_blank=False)
    dv.error = "Choisis Panama ou Cap dans la liste."
    dv.prompt = "Route empruntee par la cargaison vers l'Asie"
    ws.add_data_validation(dv)
    dv.add(cell_route)

    ws.cell(row=ligne_route + 2, column=1,
            value="Cellules en jaune = a ajuster toi-meme (voir README, section 'Fraicheur des hypotheses')")
    ws.cell(row=ligne_route + 2, column=1).font = Font(italic=True, size=9)

    return {
        "boil_off": "Hypotheses!$B$4",
        "feedgas_mult": "Hypotheses!$B$5",
        "tolling_fee": "Hypotheses!$B$6",
        "jours_europe": "Hypotheses!$B$7",
        "jours_asie_panama": "Hypotheses!$B$8",
        "jours_asie_cap": "Hypotheses!$B$9",
        "peage_panama": "Hypotheses!$B$10",
        "mwh_to_mmbtu": "Hypotheses!$B$11",
        "regaz_europe": "Hypotheses!$B$12",
        "regaz_asie": "Hypotheses!$B$13",
        "ttf": "Hypotheses!$B$17",
        "eur_usd": "Hypotheses!$B$18",
        "henry_hub": "Hypotheses!$B$19",
        "jkm": "Hypotheses!$B$20",
        "taux_fret": "Hypotheses!$B$21",
        "route": "Hypotheses!$B$22",
        "cargo_mmbtu": f"(Hypotheses!$B$3*1000000)",
    }


def _construire_feuille_netback(wb, refs):
    ws = wb.create_sheet("Netback")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 16

    ws["A1"] = "Comparaison Europe / Asie (meme formule que src/netback.py)"
    ws["A1"].font = TITRE_FONT

    _entete(ws, "B2", "Europe")
    _entete(ws, "C2", "Asie")

    lignes = [
        ("Prix rendu ($/MMBtu)",
         f"={refs['ttf']}/{refs['mwh_to_mmbtu']}*{refs['eur_usd']}",
         f"={refs['jkm']}"),
        ("Cout FOB ($/MMBtu)",
         f"={refs['feedgas_mult']}*{refs['henry_hub']}+{refs['tolling_fee']}",
         "=B4"),
        ("Jours de mer",
         f"={refs['jours_europe']}",
         f'=IF({refs["route"]}="Panama",{refs["jours_asie_panama"]},{refs["jours_asie_cap"]})'),
        ("Fret ($/MMBtu)",
         f"={refs['taux_fret']}*B5/{refs['cargo_mmbtu']}",
         f'=({refs["taux_fret"]}*C5+IF({refs["route"]}="Panama",{refs["peage_panama"]},0))/{refs["cargo_mmbtu"]}'),
        ("Perte boil-off ($/MMBtu)",
         f"={refs['boil_off']}*B5*B3",
         f"={refs['boil_off']}*C5*C3"),
        ("Frais de regaz ($/MMBtu)",
         f"={refs['regaz_europe']}",
         f"={refs['regaz_asie']}"),
    ]
    ligne = 3
    for libelle, formule_europe, formule_asie in lignes:
        ws.cell(row=ligne, column=1, value=libelle)
        ws.cell(row=ligne, column=2, value=formule_europe).number_format = "0.000"
        ws.cell(row=ligne, column=3, value=formule_asie).number_format = "0.000"
        ligne += 1
    # lignes 3 a 8 : prix_rendu(3), cout_fob(4), jours(5), fret(6), boiloff(7), regaz(8)

    ws.cell(row=9, column=1, value="Netback ($/MMBtu)").font = ENTETE_FONT
    for col, lettre in ((2, "B"), (3, "C")):
        c = ws.cell(row=9, column=col, value=f"={lettre}3-{lettre}4-{lettre}6-{lettre}7-{lettre}8")
        c.font = ENTETE_FONT
        c.number_format = "0.000"

    ws.cell(row=11, column=1, value="Destination recommandee")
    ws.cell(row=11, column=1).font = ENTETE_FONT
    ws.cell(row=11, column=2, value='=IF(B9>C9,"Europe","Asie")').font = ENTETE_FONT
    ws.cell(row=12, column=1, value="Ecart (Asie - Europe, $/MMBtu)")
    ws.cell(row=12, column=2, value="=C9-B9").number_format = "0.000"

    return ws


def _construire_feuille_sensibilite(wb, refs):
    ws = wb.create_sheet("Sensibilite")
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 14
    for col in range(3, 10):
        ws.column_dimensions[get_column_letter(col)].width = 11

    ws["A1"] = "Sensibilite : ecart de netback (Asie - Europe), en $/MMBtu"
    ws["A1"].font = TITRE_FONT
    ws["A2"] = ("Positif (vert) = l'Asie l'emporte. Negatif (rouge) = l'Europe l'emporte. "
                "La ligne ou la couleur change de signe est la frontiere de bascule.")
    ws["A2"].font = Font(italic=True, size=9)
    ws.merge_cells("A2:I2")

    ws["B3"] = "Taux de fret \\ Ecart JKM-TTF"
    ws["B3"].font = ENTETE_FONT

    ecarts = [-1.0, -0.5, 0.0, 0.5, 1.0, 1.5, 2.0]
    taux_fret_valeurs = [30_000, 50_000, 70_000, 90_000, 110_000, 130_000, 150_000]

    for i, ecart in enumerate(ecarts):
        col = 3 + i  # C..I
        cell = ws.cell(row=3, column=col, value=ecart)
        cell.number_format = "+0.00;-0.00"
        cell.font = ENTETE_FONT
        cell.alignment = Alignment(horizontal="center")

    for i, taux in enumerate(taux_fret_valeurs):
        row = 4 + i  # 4..10
        cell = ws.cell(row=row, column=2, value=taux)
        cell.number_format = "#,##0"
        cell.font = ENTETE_FONT

    ttf = refs["ttf"]
    mwh = refs["mwh_to_mmbtu"]
    eur_usd = refs["eur_usd"]
    ttf_usd_mmbtu = f"({ttf}/{mwh}*{eur_usd})"
    cargo = refs["cargo_mmbtu"]
    jours_e = refs["jours_europe"]
    jours_a = f'IF({refs["route"]}="Panama",{refs["jours_asie_panama"]},{refs["jours_asie_cap"]})'
    peage = f'IF({refs["route"]}="Panama",{refs["peage_panama"]},0)'
    fob = f"({refs['feedgas_mult']}*{refs['henry_hub']}+{refs['tolling_fee']})"
    boil = refs["boil_off"]
    regaz_e = refs["regaz_europe"]
    regaz_a = refs["regaz_asie"]

    for i in range(len(taux_fret_valeurs)):
        row = 4 + i
        taux_ref = f"$B{row}"
        for j in range(len(ecarts)):
            col = 3 + j
            lettre = get_column_letter(col)
            ecart_ref = f"{lettre}$3"

            jkm = f"({ttf_usd_mmbtu}+{ecart_ref})"
            fret_europe = f"({taux_ref}*{jours_e}/{cargo})"
            fret_asie = f"(({taux_ref}*({jours_a})+({peage}))/{cargo})"
            netback_europe = f"({ttf_usd_mmbtu}-{fob}-{fret_europe}-{boil}*{jours_e}*{ttf_usd_mmbtu}-{regaz_e})"
            netback_asie = f"({jkm}-{fob}-{fret_asie}-{boil}*({jours_a})*{jkm}-{regaz_a})"

            formule = f"={netback_asie}-{netback_europe}"
            cell = ws.cell(row=row, column=col, value=formule)
            cell.number_format = "+0.00;-0.00"

    plage = f"C4:I{3 + len(taux_fret_valeurs)}"
    regle = ColorScaleRule(
        start_type="min", start_color="F8696B",
        mid_type="num", mid_value=0, mid_color="FFFFFF",
        end_type="max", end_color="63BE7B",
    )
    ws.conditional_formatting.add(plage, regle)

    return ws


def construire_classeur(chemin_sortie=None):
    chemin_sortie = chemin_sortie or CHEMIN_SORTIE

    wb = Workbook()
    refs = _construire_feuille_hypotheses(wb)
    _construire_feuille_netback(wb, refs)
    _construire_feuille_sensibilite(wb, refs)

    chemin_sortie = Path(chemin_sortie)
    chemin_sortie.parent.mkdir(parents=True, exist_ok=True)
    wb.save(chemin_sortie)
    return chemin_sortie


if __name__ == "__main__":
    chemin = construire_classeur()
    print(f"Classeur cree : {chemin}")
