"""
Reinjecte les dernieres valeurs de marche (TTF, Henry Hub, EUR/USD) dans
le classeur Excel du Jalon 6, sans toucher aux formules ni aux
hypotheses que tu ajustes toi-meme a la main (JKM, taux de fret, route
Asie) - voir README, section "Fraicheur des hypotheses".

C'est le pont entre le pipeline Python (Jalons 4-5) et le modele de
decision Excel (Jalon 6) : lance ce script, ouvre le classeur, les
chiffres sont a jour et tout se recalcule tout seul.
"""

import logging
from pathlib import Path

from openpyxl import load_workbook

from . import data_prices

logger = logging.getLogger(__name__)

CHEMIN_CLASSEUR = Path(__file__).resolve().parent.parent / "excel" / "modele_arbitrage.xlsx"

# Les 3 seules cellules de la feuille "Hypotheses" que ce script a le
# droit d'ecrire. Tout le reste du classeur (formules, JKM, taux de
# fret, route, mise en forme, menu deroulant, heatmap) n'est JAMAIS
# touche : on ne fait que charger le classeur existant et reecrire ces
# cellules precises avant de sauvegarder.
CELLULE_TTF = "B17"
CELLULE_EUR_USD = "B18"
CELLULE_HENRY_HUB = "B19"


def _derniere_valeur(df, colonne):
    """Valeur la plus recente d'un DataFrame de prix (deja trie par date)."""
    if df is None or df.empty:
        return None
    return float(df.iloc[-1][colonne])


def recuperer_dernieres_valeurs():
    """
    Va chercher les derniers TTF, Henry Hub et EUR/USD connus (Jalon 4).
    Une valeur manquante (source en panne) devient None - elle sera
    simplement ignoree par mettre_a_jour_classeur(), qui ne casera rien.
    """
    return {
        "ttf": _derniere_valeur(data_prices.fetch_ttf(), "ttf_eur_mwh"),
        "henry_hub": _derniere_valeur(data_prices.fetch_henry_hub(), "henry_hub_usd_mmbtu"),
        "eur_usd": _derniere_valeur(data_prices.fetch_eur_usd(), "eur_usd"),
    }


def mettre_a_jour_classeur(valeurs=None, chemin_classeur=CHEMIN_CLASSEUR):
    """
    Ouvre le classeur existant et ecrit les valeurs de marche dans les
    cellules d'entree dediees - SANS recreer le classeur ni toucher aux
    formules des feuilles Netback/Sensibilite.

    Le point cle : load_workbook() est appele avec data_only=False (le
    defaut), ce qui garde le TEXTE de chaque formule ("=B3-B4-B6-B7-B8"
    par exemple) au lieu de son dernier resultat calcule. openpyxl ne
    sait pas evaluer une formule ; il se contente de lire/ecrire le
    fichier. Tant qu'on n'assigne pas explicitement de nouvelle valeur a
    une cellule, sa formule est reecrite a l'identique lors du wb.save().
    C'est pour ca qu'on ne touche QUE les 3 cellules listees plus haut.

    Consequence a connaitre : openpyxl ne reecrit pas non plus le
    resultat mis en cache des formules qu'il ne modifie pas. Si le
    classeur etait configure en recalcul "Manuel", Excel pourrait
    afficher d'anciens chiffres jusqu'a un F9. On force donc un
    recalcul complet a l'ouverture (fullCalcOnLoad), pour que "le
    classeur se recalcule seul" soit garanti, quel que soit le mode de
    calcul enregistre dans le fichier.
    """
    valeurs = valeurs if valeurs is not None else recuperer_dernieres_valeurs()

    wb = load_workbook(chemin_classeur)
    ws = wb["Hypotheses"]

    ecrit = {}
    correspondance = {
        "ttf": CELLULE_TTF,
        "eur_usd": CELLULE_EUR_USD,
        "henry_hub": CELLULE_HENRY_HUB,
    }
    for cle, cellule in correspondance.items():
        valeur = valeurs.get(cle)
        if valeur is not None:
            ws[cellule] = valeur
            ecrit[cle] = valeur
        else:
            logger.warning("Valeur '%s' indisponible : cellule %s laissee inchangee.", cle, cellule)

    wb.calculation.fullCalcOnLoad = True
    wb.save(chemin_classeur)
    return ecrit


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    ecrit = mettre_a_jour_classeur()
    if ecrit:
        print("Classeur mis a jour :")
        for cle, valeur in ecrit.items():
            print(f"  {cle}: {valeur}")
    else:
        print("Aucune valeur recuperee : classeur inchange.")
