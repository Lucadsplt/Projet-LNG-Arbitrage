import pytest
from openpyxl import load_workbook

from src.build_excel_model import construire_classeur
from src.update_excel import CELLULE_EUR_USD, CELLULE_HENRY_HUB, CELLULE_TTF, mettre_a_jour_classeur


@pytest.fixture
def classeur_temporaire(tmp_path):
    """Un classeur genere fraichement (Jalon 6), isole dans un dossier temporaire."""
    chemin = tmp_path / "modele_test.xlsx"
    construire_classeur(chemin_sortie=chemin)
    return chemin


def test_mettre_a_jour_ecrit_les_3_cellules_visees(classeur_temporaire):
    mettre_a_jour_classeur(
        valeurs={"ttf": 40.0, "eur_usd": 1.05, "henry_hub": 3.50},
        chemin_classeur=classeur_temporaire,
    )

    ws = load_workbook(classeur_temporaire)["Hypotheses"]
    assert ws[CELLULE_TTF].value == 40.0
    assert ws[CELLULE_EUR_USD].value == 1.05
    assert ws[CELLULE_HENRY_HUB].value == 3.50


def test_mettre_a_jour_ne_casse_aucune_formule(classeur_temporaire):
    # On releve toutes les formules du classeur AVANT la mise a jour...
    wb_avant = load_workbook(classeur_temporaire)
    formules_avant = {}
    for nom_feuille in wb_avant.sheetnames:
        ws = wb_avant[nom_feuille]
        for ligne in ws.iter_rows():
            for cellule in ligne:
                if isinstance(cellule.value, str) and cellule.value.startswith("="):
                    formules_avant[(nom_feuille, cellule.coordinate)] = cellule.value

    assert len(formules_avant) > 50  # le classeur du Jalon 6 est riche en formules

    mettre_a_jour_classeur(
        valeurs={"ttf": 40.0, "eur_usd": 1.05, "henry_hub": 3.50},
        chemin_classeur=classeur_temporaire,
    )

    # ...et on verifie qu'elles sont TOUTES identiques apres.
    wb_apres = load_workbook(classeur_temporaire)
    for (nom_feuille, coord), formule_avant in formules_avant.items():
        assert wb_apres[nom_feuille][coord].value == formule_avant


def test_mettre_a_jour_ignore_une_valeur_manquante(classeur_temporaire):
    ws_avant = load_workbook(classeur_temporaire)["Hypotheses"]
    ttf_avant = ws_avant[CELLULE_TTF].value
    henry_hub_avant = ws_avant[CELLULE_HENRY_HUB].value

    mettre_a_jour_classeur(
        valeurs={"ttf": None, "eur_usd": 1.10, "henry_hub": None},
        chemin_classeur=classeur_temporaire,
    )

    ws = load_workbook(classeur_temporaire)["Hypotheses"]
    assert ws[CELLULE_TTF].value == ttf_avant  # inchangee : pas de valeur fournie
    assert ws[CELLULE_EUR_USD].value == 1.10
    assert ws[CELLULE_HENRY_HUB].value == henry_hub_avant  # inchangee : pas de valeur fournie


def test_mettre_a_jour_preserve_le_menu_deroulant_et_la_heatmap(classeur_temporaire):
    mettre_a_jour_classeur(
        valeurs={"ttf": 40.0, "eur_usd": 1.05, "henry_hub": 3.50},
        chemin_classeur=classeur_temporaire,
    )

    wb = load_workbook(classeur_temporaire)
    ws_hyp = wb["Hypotheses"]
    assert len(ws_hyp.data_validations.dataValidation) == 1
    assert ws_hyp.data_validations.dataValidation[0].formula1 == '"Panama,Cap"'

    ws_sens = wb["Sensibilite"]
    assert len(ws_sens.conditional_formatting._cf_rules) >= 1
